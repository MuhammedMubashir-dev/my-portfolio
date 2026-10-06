"""Build the two evidence-based mobile resumes reviewed on 1 October 2026."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "resume" / "2026-10-06"
PORTFOLIO = "https://muhammed-mubashir-portfolio.netlify.app"
BLACK = RGBColor(0, 0, 0)

RESUMES = {
    "Flutter": {
        "headline": "Flutter Developer",
        "summary": (
            "Flutter developer with company experience building and maintaining POS, customer ordering, "
            "and delivery applications using Dart, Provider, and REST APIs. Contributions include cash "
            "management, bilingual printing, payments, reporting, and regression tests. Additional "
            "React Native experience covers Android and iOS authentication and deep linking."
        ),
        "skills": [
            ("Mobile", "Flutter, Dart, Provider, React Native, Android and iOS, responsive UI"),
            ("Integrations", "REST APIs, Razorpay, Google Maps, PDF and USB receipt printing, Arabic/RTL localization"),
            ("Engineering", "Git, GitHub, Postman, unit and widget testing"),
            ("Supporting web", "React, Next.js, JavaScript, TypeScript, Tailwind CSS, Stripe"),
        ],
        "experience": [
            "Develop and maintain features in existing Flutter codebases for POS, laundry ordering, and staff delivery apps.",
            "Coordinate API contracts with backend developers and investigate request, payload, and authentication issues.",
            "Verify changes with unit, widget, and regression tests, and support Android and iOS release configuration.",
        ],
        "trainee": [
            "Implemented Flutter interfaces and REST API integration for pickup and delivery workflows.",
        ],
        "projects": [
            {
                "name": "EPOSMOB",
                "stack": "Retail POS and Inventory | Flutter, Provider",
                "link": f"{PORTFOLIO}/#project-5",
                "bullets": [
                    "Built shift opening and Day Close workflows with cash-denomination calculations and pending-close alerts.",
                    "Implemented purchase-return, packing, and shipping workflows.",
                    "Developed reporting and transaction listings with shared filters and validated Excel exports.",
                    "Fixed English/Arabic invoice rendering, USB printing, barcode pricing, and stock isolation, with regression tests.",
                ],
            },
            {
                "name": "Ganvin Customer",
                "stack": "Laundry Ordering and Tracking | Flutter, Provider",
                "bullets": [
                    "Integrated Razorpay checkout and corrected payment-confirmation messaging.",
                    "Reduced My Orders API requests from six to two by reusing responses for lists and counts.",
                    "Added update prompts and maintenance-mode handling, and improved cart, login navigation, and saved addresses.",
                ],
            },
            {
                "name": "Ganvin Executive",
                "stack": "Staff Pickup and Delivery | Flutter",
                "link": f"{PORTFOLIO}/#project-2",
                "bullets": [
                    "Implemented independent pagination and date filters across four delivery stages.",
                    "Added Google Maps navigation and saved locations, and improved OTP entry and empty/error states.",
                ],
            },
            {
                "name": "Connect App",
                "stack": "Business Networking | React Native, TypeScript",
                "bullets": [
                    "Implemented LinkedIn login on Android and iOS with PKCE and native modules, and fixed deferred deep links.",
                ],
            },
        ],
        "web": (
            "CloudPOS and Juice World: contributed Next.js storefront themes, tenant-aware Stripe checkout, "
            "Razorpay integration, English/Arabic localization, ISR, and SEO metadata."
        ),
    },
    "Cross-Platform": {
        "headline": "Cross Platform Mobile Developer",
        "summary": (
            "Mobile developer with company experience delivering Flutter and React Native features for "
            "POS, logistics, and business networking applications on Android and iOS. Skilled in API "
            "integration, state management, payments, authentication, and deep linking. Also contributes "
            "React and Next.js storefront features, checkout integrations, and localization."
        ),
        "skills": [
            ("Mobile", "Flutter, Dart, Provider, ChangeNotifier, setState, React Native, TypeScript, JavaScript"),
            ("Integrations", "REST APIs, JSON, OAuth with PKCE, Branch deep linking, Razorpay, Google Maps"),
            ("Engineering", "Git, GitHub, Postman, unit and widget tests, responsive UI, Android and iOS builds"),
            ("Supporting web", "React, Next.js, Tailwind CSS, Stripe, ISR, English/Arabic i18n and RTL"),
        ],
        "experience": [
            "Deliver Flutter and React Native features in existing company codebases; coordinate API "
            "contracts with backend developers and resolve mobile authentication, data, and release issues."
        ],
        "projects": [
            {
                "name": "Connect App",
                "stack": "Business Networking | React Native, TypeScript",
                "link": f"{PORTFOLIO}/#project-4",
                "bullets": [
                    "Implemented LinkedIn authentication on Android and iOS with PKCE, native modules, browser callbacks, and backend API integration; added authentication regression tests.",
                    "Resolved Branch deferred deep-link startup races and navigation after login; integrated API-provided profile sharing URLs and Android/iOS release configuration fixes.",
                    "Developed company profiles, connection and subscription workflows; integrated reCAPTCHA Enterprise and API-driven social platforms, and improved profile-image handling.",
                ],
            },
            {
                "name": "EPOSMOB",
                "stack": "Retail POS and Inventory | Flutter, Provider",
                "link": f"{PORTFOLIO}/#project-5",
                "bullets": [
                    "Built shift opening, Day Close, purchase-return, and order packing/shipping workflows; implemented bilingual receipt printing and Arabic/RTL layouts.",
                    "Developed shared listing filters and validated Excel exports; corrected pricing and stock isolation, guarded stale invoice requests, and added regression tests.",
                ],
            },
            {
                "name": "Ganvin Customer",
                "stack": "Laundry Ordering and Tracking | Flutter, Provider",
                "bullets": [
                    "Integrated Razorpay and payment-confirmation handling; reduced My Orders API requests from six to two and added update prompts and maintenance handling.",
                ],
            },
            {
                "name": "Ganvin Executive",
                "stack": "Staff Pickup and Delivery | Flutter, setState",
                "bullets": [
                    "Implemented independent pagination across four delivery stages, API date filters, and Google Maps navigation; improved OTP entry and loading/error handling.",
                ],
            },
        ],
        "web": (
            "CloudPOS and Juice World: built Next.js storefront themes and English/Arabic interfaces; "
            "implemented tenant-aware Stripe checkout, Razorpay integration, ISR, and SEO metadata."
        ),
    },
}


def hyperlink(paragraph, label, url, *, size=10.5, bold=False):
    relation = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    font = OxmlElement("w:rFonts")
    font.set(qn("w:ascii"), "Calibri")
    font.set(qn("w:hAnsi"), "Calibri")
    properties.append(font)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    properties.append(color)
    point_size = OxmlElement("w:sz")
    point_size.set(qn("w:val"), str(round(size * 2)))
    properties.append(point_size)
    if bold:
        properties.append(OxmlElement("w:b"))
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def paragraph(doc, text="", *, bold=False, after=2.5, before=0, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    p.add_run(text).bold = bold
    return p


def bullet(doc, text):
    p = paragraph(doc, "\u2022 " + text, after=3)
    p.paragraph_format.left_indent = Inches(0.13)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    return p


def heading(doc, text):
    return doc.add_paragraph(text, style="Heading 1")


def build_resume(variant, content):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.53)
    section.left_margin = section.right_margin = Inches(0.62)

    for name in ("Normal", "Title", "Subtitle", "Heading 1"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(2.5)
        style.paragraph_format.line_spacing = 1.03
        style.paragraph_format.widow_control = True
    doc.styles["Title"].font.size = Pt(21)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].paragraph_format.space_after = Pt(2)
    doc.styles["Subtitle"].font.size = Pt(12)
    doc.styles["Subtitle"].font.bold = True
    doc.styles["Subtitle"].paragraph_format.space_after = Pt(4)
    doc.styles["Heading 1"].font.size = Pt(11)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(7)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(3)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True

    doc.add_paragraph("MUHAMMED MUBASHIR K", style="Title")
    doc.add_paragraph(content["headline"], style="Subtitle")
    p = paragraph(doc, "Malappuram, Kerala, India | ")
    hyperlink(p, "+91 8089433955", "tel:+918089433955")
    p.add_run(" | ")
    hyperlink(p, "muhammedmubashir720@gmail.com", "mailto:muhammedmubashir720@gmail.com")
    p = paragraph(doc)
    for index, (label, url) in enumerate([
        ("muhammed-mubashir-portfolio.netlify.app", PORTFOLIO),
        ("github.com/MuhammedMubashir-dev", "https://github.com/MuhammedMubashir-dev"),
        ("linkedin.com/in/muhammed-mubashir-k", "https://www.linkedin.com/in/muhammed-mubashir-k"),
    ]):
        if index:
            p.add_run(" | ").font.size = Pt(9.5)
        hyperlink(p, label, url, size=9.5)

    heading(doc, "Professional Summary")
    paragraph(doc, content["summary"])
    heading(doc, "Technical Skills")
    for label, value in content["skills"]:
        p = paragraph(doc)
        p.add_run(label + ": ").bold = True
        p.add_run(value)

    heading(doc, "Work Experience")
    paragraph(doc, "ENKE Consulting Services LLP | Apr 2026 \u2013 Present", bold=True, keep=True)
    paragraph(doc, "Junior Application Developer | Jul 2026 \u2013 Present", bold=True, keep=True)
    for text in content["experience"]:
        bullet(doc, text)
    paragraph(doc, "Full Stack Developer Trainee | Apr 2026 \u2013 Jul 2026", before=2, keep=True)
    for text in content.get("trainee", []):
        bullet(doc, text)

    heading(doc, "Selected Company Projects")
    for project in content["projects"]:
        p = paragraph(doc, before=3.5, after=2, keep=True)
        if project.get("link"):
            hyperlink(p, project["name"], project["link"], bold=True)
        else:
            p.add_run(project["name"]).bold = True
        p.add_run(" | " + project["stack"])
        for text in project["bullets"]:
            bullet(doc, text)

    heading(doc, "Supporting Web Experience")
    paragraph(doc, content["web"])
    heading(doc, "Education")
    paragraph(doc, "Bachelor of Computer Applications | 2023 \u2013 2026", bold=True, keep=True)
    paragraph(doc, "Priyadarshini Arts and Science College, Melmuri, Malappuram | University of Calicut", after=0)

    # Remove inherited decorative borders and theme colors from Word templates.
    for root in (doc.styles.element, doc.element):
        for border in root.xpath(".//w:pBdr"):
            border.getparent().remove(border)
        for color in root.xpath(".//w:color"):
            for attribute in ("themeColor", "themeTint", "themeShade"):
                color.attrib.pop(qn("w:" + attribute), None)
            color.set(qn("w:val"), "000000")
    # Thin rule under each section heading, added after template borders are stripped.
    for p in doc.paragraphs:
        if p.style.name == "Heading 1":
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for key, value in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "000000")):
                bottom.set(qn("w:" + key), value)
            borders.append(bottom)
            p._p.get_or_add_pPr().append(borders)

    doc.core_properties.title = "Muhammed Mubashir K " + content["headline"] + " Resume"
    doc.core_properties.subject = "Company mobile development experience"
    doc.core_properties.author = "Muhammed Mubashir K"
    doc.core_properties.keywords = "Flutter, React Native, Dart, TypeScript, mobile development"
    doc.core_properties.comments = ""
    path = OUTPUT / f"Muhammed-Mubashir-{variant}-Resume.docx"
    doc.save(path)
    print(path)


if __name__ == "__main__":
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for variant, content in RESUMES.items():
        build_resume(variant, content)
