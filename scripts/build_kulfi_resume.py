"""Build the résumé tailored to Kulfi's Software Engineer I (Product / Full-Stack) role.

Produces a DOCX (python-docx) and a PDF (reportlab) from one content model so the
two stay identical. Run from anywhere:

    python scripts/build_kulfi_resume.py
"""
from html import escape
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor, black
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer

from build_resume import add_hyperlink

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "resume"
OUT.mkdir(parents=True, exist_ok=True)
STEM = "Muhammed-Mubashir-Resume-Kulfi"

NAME = "MUHAMMED MUBASHIR K"
HEADLINE = "Junior Full Stack Developer  |  Flutter, React Native, Node.js"
CONTACT = [
    ("Malappuram, Kerala, India (open to relocating to Bengaluru)", None),
    ("+91 8089433955", "tel:+918089433955"),
    ("muhammedmubashir720@gmail.com", "mailto:muhammedmubashir720@gmail.com"),
]
LINKS = [
    ("Portfolio", "https://muhammed-mubashir-portfolio.netlify.app"),
    ("GitHub", "https://github.com/MuhammedMubashir-dev"),
    ("LinkedIn", "https://www.linkedin.com/in/muhammed-mubashir-k"),
]

SUMMARY = (
    "Application developer who ships product features end to end. At ENKE Consulting I have worked on "
    "six production apps since April 2026, with 71 merged pull requests across Flutter, React Native, and "
    "Next.js codebases. Outside work I built a three-app laundry marketplace from the Flutter clients to the "
    "Express and PostgreSQL API. I handle the loading, empty, error, and retry states, write regression "
    "tests before handoff, and use Claude Code and Codex with every line reviewed before it ships."
)

# Each project: (title line, [bullets], optional (label, url) link)
SHIPPED = [
    (
        "EPOSMOB  |  Flutter point of sale  |  52 merged PRs since July 2026",
        [
            "Consolidated 25+ list and report screens onto one shared listing UI with validated Excel export "
            "that fetches every page and rejects duplicate or incomplete results.",
            "Built Day Close with cash-denomination validation, bilingual English/Arabic receipts and PDF "
            "invoices, and regression tests for pagination, stale responses, and filter state.",
        ],
        ("EPOSMOB case study", "https://muhammed-mubashir-portfolio.netlify.app/#project-5"),
    ),
    (
        "Ganvin Customer and Executive  |  Flutter delivery apps",
        [
            "Integrated Razorpay checkout with correct confirmation states, plus optional and mandatory "
            "update prompts and a maintenance mode; cut order-history API calls from six to two by reusing responses.",
            "Implemented independent pagination and date filters per delivery stage, fixed OTP and endless-loading "
            "failures, and traced missing orders to a server timezone mismatch with the backend team.",
        ],
        ("Ganvin case study", "https://muhammed-mubashir-portfolio.netlify.app/#project-6"),
    ),
    (
        "Connect App  |  React Native, Kotlin, Swift",
        [
            "Implemented LinkedIn sign-in with PKCE and native modules on Android and iOS, kept deferred deep links "
            "alive through app startup, and added authentication regression tests.",
        ],
        None,
    ),
    (
        "CloudPOS and Juice World  |  Next.js storefronts",
        [
            "Delivered tenant-aware Stripe and Razorpay checkout, reusable theme packs, and an English/Arabic "
            "right-to-left storefront with ISR and SEO metadata.",
        ],
        None,
    ),
]

PERSONAL = [
    (
        "Laundry pickup and delivery marketplace  |  Flutter, Express, PostgreSQL  |  2026, private repos, demo on request",
        [
            "Three Flutter apps (customer, rider, shop) in one Dart workspace sharing a core package: dio client with "
            "rotating refresh tokens, Socket.io live tracking, OTP login, and Firebase push wired behind config.",
            "Express 5 API on PostgreSQL with PostGIS nearby-shop search, pickup slots, rider assignment, and an "
            "end-to-end order-flow test; React and Vite admin panel for shops, riders, and orders.",
        ],
        None,
    ),
    (
        "StockFlow  |  Offline-first Flutter inventory app  |  public, GitHub Actions CI, Android release",
        [
            "Riverpod state with repository, domain, and presentation layers; responsive Material 3 for phone, "
            "tablet, and web; English/Arabic RTL; light and dark themes; widget tests on the dashboard flow.",
        ],
        ("github.com/MuhammedMubashir-dev/stockflow-flutter", "https://github.com/MuhammedMubashir-dev/stockflow-flutter"),
    ),
    (
        "Next.js e-commerce store  |  live on Vercel",
        [
            "App Router storefront with ISR, Algolia typo-tolerant search, category filters, persistent cart and "
            "checkout, and React Hook Form with Zod validation.",
        ],
        ("nextjs-ecommerce-amber.vercel.app", "https://nextjs-ecommerce-amber.vercel.app/"),
    ),
]

SKILLS = [
    ("Mobile", "Flutter, Dart, Provider, Riverpod, React Native, TypeScript, Android and iOS releases"),
    ("Backend and data", "Node.js, Express, PostgreSQL and PostGIS, REST APIs, JWT, Socket.io, Firebase"),
    ("Web", "React, Next.js, Tailwind CSS, ISR, SEO"),
    ("Practice", "Unit, widget, and regression tests; GitHub Actions CI; Postman; Git; Claude Code and Codex with every change reviewed"),
]

EXPERIENCE = [
    (
        "ENKE Consulting Services LLP, Kerala  |  Apr 2026 - Present",
        [
            "Junior Full Stack Developer (Jul 2026 - Present, permanent from Oct 2026): feature development across POS, delivery, and "
            "commerce products; API contracts agreed with backend developers before building screens.",
            "Full Stack Developer Trainee (Apr - Jul 2026): Flutter interfaces and REST integration for pickup and "
            "delivery workflows; shared React and Next.js storefront components. Promoted after three months.",
        ],
    ),
]

EDUCATION = (
    "Bachelor of Computer Applications  |  2023 - 2026",
    "Priyadarshini Arts and Science College, Melmuri, Malappuram  |  University of Calicut",
)


# ----------------------------------------------------------------------------- DOCX
def build_docx(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(0.45)
    sec.left_margin = sec.right_margin = Inches(0.6)
    for name in ["Normal", "Title", "Heading 1", "List Bullet"]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(9.5)
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_after = Pt(2)
        st.paragraph_format.line_spacing = 1.04
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.size = Pt(10.5)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(6)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True

    def para(text, bold=False):
        p = doc.add_paragraph()
        p.add_run(text).bold = bold
        return p

    def bullet(text):
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.14)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        return p

    def heading(text):
        return doc.add_paragraph(text, style="Heading 1")

    def project(title, bullets, link):
        para(title, True)
        for b in bullets:
            bullet(b)
        if link:
            p = doc.add_paragraph()
            add_hyperlink(p, link[0], link[1], size=9.5)

    doc.add_paragraph(NAME, style="Title")
    para(HEADLINE, True)
    p = doc.add_paragraph()
    for i, (text, url) in enumerate(CONTACT):
        if i:
            p.add_run("  |  ")
        if url:
            add_hyperlink(p, text, url, size=9.5)
        else:
            p.add_run(text)
    p = doc.add_paragraph()
    for i, (label, url) in enumerate(LINKS):
        if i:
            p.add_run("  |  ")
        add_hyperlink(p, label, url, size=9.5)

    heading("Summary")
    para(SUMMARY)

    heading("Shipped at ENKE Consulting (production, since April 2026)")
    for title, bullets, link in SHIPPED:
        project(title, bullets, link)

    heading("Personal projects (built end to end)")
    for title, bullets, link in PERSONAL:
        project(title, bullets, link)

    heading("Skills")
    for label, value in SKILLS:
        p = doc.add_paragraph()
        p.add_run(label + ": ").bold = True
        p.add_run(value)

    heading("Experience")
    for title, bullets in EXPERIENCE:
        para(title, True)
        for b in bullets:
            bullet(b)

    heading("Education")
    para(EDUCATION[0], True)
    para(EDUCATION[1])

    doc.core_properties.title = "Muhammed Mubashir K Resume"
    doc.core_properties.author = "Muhammed Mubashir K"
    for element in [doc.styles.element, doc.element]:
        for border in element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    doc.save(path)


# ------------------------------------------------------------------------------ PDF
LINK = HexColor("#1a4fa3")


def pdf_styles():
    base = dict(fontName="Helvetica", fontSize=9.2, leading=11.1, textColor=black)
    return {
        "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=18, leading=20, alignment=TA_CENTER, spaceAfter=1),
        "headline": ParagraphStyle("headline", fontName="Helvetica-Bold", fontSize=9.8, leading=12, alignment=TA_CENTER, spaceAfter=1),
        "contact": ParagraphStyle("contact", alignment=TA_CENTER, spaceAfter=1, **base),
        "h": ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=10, leading=12, spaceBefore=5.5, spaceAfter=1.5, textColor=black),
        "p": ParagraphStyle("p", spaceAfter=2, **base),
        "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=9.2, leading=11.1, spaceBefore=2.5, spaceAfter=0.5, textColor=black),
        "b": ParagraphStyle("b", leftIndent=11, bulletIndent=2, spaceAfter=1.5, **base),
        "link": ParagraphStyle("link", spaceAfter=2, textColor=LINK, **{k: v for k, v in base.items() if k != "textColor"}),
    }


def a(text, url):
    return f'<link href="{escape(url)}" color="#1a4fa3">{escape(text)}</link>'


def build_pdf(path):
    s = pdf_styles()
    doc = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=0.6 * inch, rightMargin=0.6 * inch,
        topMargin=0.45 * inch, bottomMargin=0.45 * inch,
        title="Muhammed Mubashir K Resume", author="Muhammed Mubashir K",
    )
    story = [
        Paragraph(escape(NAME), s["name"]),
        Paragraph(escape(HEADLINE), s["headline"]),
        Paragraph("  |  ".join(a(t, u) if u else escape(t) for t, u in CONTACT), s["contact"]),
        Paragraph("  |  ".join(a(t, u) for t, u in LINKS), s["contact"]),
        Spacer(1, 2),
        Paragraph("Summary", s["h"]),
        Paragraph(escape(SUMMARY), s["p"]),
    ]

    def project(title, bullets, link):
        block = [Paragraph(escape(title), s["title"])]
        block += [Paragraph(escape(b), s["b"], bulletText="•") for b in bullets]
        if link:
            block.append(Paragraph(a(link[0], link[1]), s["link"]))
        return KeepTogether(block)

    story.append(Paragraph("Shipped at ENKE Consulting (production, since April 2026)", s["h"]))
    story += [project(*item) for item in SHIPPED]
    story.append(Paragraph("Personal projects (built end to end)", s["h"]))
    story += [project(*item) for item in PERSONAL]
    story.append(Paragraph("Skills", s["h"]))
    story += [Paragraph(f"<b>{escape(label)}:</b> {escape(value)}", s["p"]) for label, value in SKILLS]
    story.append(Paragraph("Experience", s["h"]))
    for title, bullets in EXPERIENCE:
        story.append(KeepTogether([Paragraph(escape(title), s["title"])] + [Paragraph(escape(b), s["b"], bulletText="•") for b in bullets]))
    story.append(Paragraph("Education", s["h"]))
    story.append(Paragraph(escape(EDUCATION[0]), s["title"]))
    story.append(Paragraph(escape(EDUCATION[1]), s["p"]))
    doc.build(story)
    return doc.page


if __name__ == "__main__":
    docx_path = OUT / f"{STEM}.docx"
    pdf_path = OUT / f"{STEM}.pdf"
    build_docx(docx_path)
    pages = build_pdf(pdf_path)
    print(docx_path)
    print(pdf_path, "pages:", pages)
