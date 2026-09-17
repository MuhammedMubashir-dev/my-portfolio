from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from build_resume import add_hyperlink

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'resume'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.58)
sec.left_margin = sec.right_margin = Inches(.65)
for name in ['Normal', 'Title', 'Heading 1', 'List Bullet']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.05
doc.styles['Title'].font.size = Pt(22)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(11)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True

def para(text, bold=False):
    p = doc.add_paragraph()
    p.add_run(text).bold = bold
    return p

def bullet(text):
    p = doc.add_paragraph(text, style='List Bullet')
    p.paragraph_format.left_indent = Inches(.14)
    p.paragraph_format.first_line_indent = Inches(-.14)
    return p

def heading(text):
    return doc.add_paragraph(text, style='Heading 1')

doc.add_paragraph('MUHAMMED MUBASHIR K', style='Title')
para('Flutter Developer | Mobile & POS Applications', True)
p = para('Kerala, India | ')
add_hyperlink(p, '+91 8089433955', 'tel:+918089433955', size=10)
p.add_run(' | ')
add_hyperlink(p, 'muhammedmubashir720@gmail.com', 'mailto:muhammedmubashir720@gmail.com', size=10)
p = doc.add_paragraph()
for i, (label, url) in enumerate([
    ('Portfolio', 'https://muhammed-mubashir-portfolio.netlify.app'),
    ('GitHub', 'https://github.com/MuhammedMubashir-dev'),
    ('LinkedIn', 'https://www.linkedin.com/in/muhammed-mubashir-k'),
]):
    if i: p.add_run(' | ')
    add_hyperlink(p, label, url, size=10)

heading('Professional Summary')
para('Flutter developer with production experience building POS, customer ordering, and delivery applications using Dart, Provider, and REST APIs. Delivered cash-management workflows, payment integrations, bilingual printing, and API optimizations, with regression tests for pricing and stock handling. Promoted to Junior Application Developer after three months at ENKE Consulting.')

heading('Technical Skills')
for label, value in [
    ('Flutter and Dart', 'Provider, ChangeNotifier, setState, reusable widgets, responsive layouts'),
    ('Integrations', 'REST APIs, JSON, Razorpay, Stripe, Google Maps, push notifications, receipt printing, RTL'),
    ('Web Development', 'React, Next.js, JavaScript, TypeScript, Tailwind CSS'),
    ('Engineering Tools', 'Git, GitHub, Postman, VS Code, SQLite, SharedPreferences'),
]:
    p = doc.add_paragraph()
    p.add_run(label + ': ').bold = True
    p.add_run(value)

heading('Work Experience')
para('ENKE Consulting Services LLP | Apr 2026 - Present', True)
para('Junior Application Developer | Jul 2026 - Present', True)
bullet('Delivered production Flutter features across POS and logistics applications; diagnosed API edge cases, authentication failures, and Android and iOS release issues.')
para('Full Stack Developer Trainee | Apr 2026 - Jul 2026', True)
bullet('Integrated REST APIs and built reusable Flutter interfaces; contributed React and Next.js commerce features and collaborated with backend developers on API issues.')

heading('Selected Projects')
para('EPOSMOB | Retail POS and Inventory | Flutter, Provider', True)
bullet('Built shift-opening and Day Close workflows with cash-denomination calculations and pending-close alerts; implemented purchase-return screens using Provider and REST APIs.')
bullet('Fixed English/Arabic invoice rendering and USB printing; corrected barcode pricing and variant-specific stock selection, adding regression tests for pricing fallbacks and stock isolation.')
p = doc.add_paragraph()
add_hyperlink(p, 'EPOSMOB case study', 'https://muhammed-mubashir-portfolio.netlify.app/#project-5', size=10)
para('Ganvin Customer | Delivery Booking and Tracking | Flutter, Provider', True)
bullet('Integrated Razorpay checkout and corrected payment status messaging so customers see a confirming state before success; improved cart responsiveness and authentication navigation.')
bullet('Reduced My Orders API calls from six to two by reusing responses for lists and counts; added optional and mandatory app-update prompts and maintenance handling.')
para('Ganvin Executive | Staff Pickup and Delivery App | Flutter, setState', True)
bullet('Implemented independent pagination and API date filters across four pickup/delivery stages; integrated Google Maps and saved locations for staff navigation.')
bullet('Replaced six OTP fields with one validated input; fixed null-response OTP failures and endless loading on empty/error list responses, and kept Verify controls visible for long item names.')
p = doc.add_paragraph()
add_hyperlink(p, 'Ganvin Executive case study', 'https://muhammed-mubashir-portfolio.netlify.app/#project-2', size=10)

heading('Education')
para('Bachelor of Computer Applications | 2023 - 2026', True)
para('Priyadarshini Arts and Science College, Melmuri, Malappuram | University of Calicut')
doc.core_properties.title = 'Muhammed Mubashir K Flutter Developer Resume'
doc.core_properties.author = 'Muhammed Mubashir K'
for element in [doc.styles.element, doc.element]:
    for border in element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
path = OUT / 'Muhammed-Mubashir-Flutter-Resume.docx'
doc.save(path)
print(path)
